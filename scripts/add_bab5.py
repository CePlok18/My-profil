from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

SOURCE = "TA ADAM REVISI BAB 4.docx"
OUTPUT = "TA ADAM REVISI BAB 4-5.docx"

doc = Document(SOURCE)

insert_before = None
for paragraph in doc.paragraphs:
    if paragraph.text.strip().upper() == "DAFTAR PUSTAKA":
        insert_before = paragraph
        break

if insert_before is None:
    raise SystemExit("DAFTAR PUSTAKA tidak ditemukan")

if any("BAB V KESIMPULAN DAN SARAN" in p.text.upper() for p in doc.paragraphs):
    raise SystemExit("BAB V sudah ada, proses dibatalkan agar tidak duplikat")


def set_font(run, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)
    run.bold = bold


def para(text="", bold=False, centered=False, first_line=True):
    p = insert_before.insert_paragraph_before()
    if text:
        for index, part in enumerate(text.split("\n")):
            if index:
                p.add_run().add_break()
            run = p.add_run(part)
            set_font(run, bold)
    if centered:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if first_line:
        p.paragraph_format.first_line_indent = Inches(0.3)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    return p


def heading(text, centered=False):
    return para(text, bold=True, centered=centered, first_line=False)


insert_before.insert_paragraph_before().add_run().add_break(WD_BREAK.PAGE)

heading("BAB V KESIMPULAN DAN SARAN", centered=True)

heading("5.1 Kesimpulan")
para(
    "Berdasarkan hasil perancangan, pembuatan, dan pengujian sistem smart loker penyimpanan berbasis Internet of Things, dapat disimpulkan bahwa sistem yang dibuat mampu mengintegrasikan aplikasi web, database, komunikasi MQTT, dan perangkat keras ESP32 dalam satu sistem yang saling terhubung. Sistem web berfungsi sebagai antarmuka untuk admin dan user, sedangkan ESP32 berfungsi sebagai pengendali perangkat keras berupa servo dan LED indikator pada setiap loker."
)
para(
    "Sistem login yang dibuat dapat membedakan hak akses antara admin dan user. Admin memiliki akses untuk memantau seluruh loker, menambahkan akun user, menghapus akun user, serta menentukan loker yang dapat digunakan oleh masing-masing user. Sementara itu, user hanya dapat mengakses satu loker yang telah diberikan oleh admin. Dengan adanya pembatasan akses tersebut, penggunaan loker menjadi lebih teratur dan aman."
)
para(
    "Penggunaan Supabase sebagai database dan autentikasi dapat membantu pengelolaan data pengguna, data loker, hak akses, riwayat aktivitas, dan perintah loker. Setiap user yang dibuat oleh admin akan disimpan pada sistem autentikasi dan dihubungkan dengan satu loker melalui tabel locker_access. Hal ini memastikan bahwa setiap user hanya memiliki akses terhadap loker yang telah ditentukan."
)
para(
    "Komunikasi antara web dan perangkat ESP32 menggunakan MQTT dapat berjalan sesuai rancangan. Ketika tombol Buka atau Tutup pada web ditekan, sistem mengirimkan command ke topic MQTT sesuai nomor loker. ESP32 kemudian menerima command tersebut, menggerakkan servo pada loker yang sesuai, serta mengubah warna LED indikator. LED berwarna hijau ketika loker dibuka dan kembali merah ketika loker ditutup."
)
para(
    "Berdasarkan hasil pengujian, fungsi utama sistem seperti login, manajemen user, pembatasan akses loker, pengiriman command, pengendalian servo, dan indikator LED dapat berjalan sesuai kebutuhan. Dengan demikian, sistem smart loker penyimpanan berbasis IoT ini dapat digunakan sebagai solusi untuk meningkatkan efisiensi dan keamanan dalam proses penyimpanan barang."
)

heading("5.2 Saran")
para(
    "Sistem smart loker penyimpanan yang telah dibuat masih dapat dikembangkan lebih lanjut agar memiliki fitur yang lebih lengkap dan andal. Pengembangan pertama yang dapat dilakukan adalah menambahkan sensor pintu atau magnetic door sensor pada setiap loker. Sensor tersebut dapat digunakan untuk mendeteksi kondisi pintu loker secara langsung, sehingga sistem tidak hanya mengetahui perintah buka dan tutup, tetapi juga mengetahui kondisi fisik pintu apakah benar-benar terbuka atau tertutup."
)
para(
    "Pengembangan berikutnya adalah menambahkan fitur notifikasi kepada admin maupun user. Notifikasi dapat diberikan ketika loker berhasil dibuka, loker belum ditutup dalam waktu tertentu, atau terjadi gangguan koneksi pada perangkat ESP32. Dengan adanya notifikasi, pengguna dan admin dapat mengetahui kondisi loker secara lebih cepat."
)
para(
    "Sistem juga dapat dikembangkan dengan menambahkan fitur riwayat penggunaan yang lebih detail. Riwayat tersebut dapat mencatat nama user, nomor loker, waktu buka, waktu tutup, serta status keberhasilan command. Data riwayat yang lengkap dapat membantu admin dalam melakukan pemantauan dan evaluasi penggunaan loker."
)
para(
    "Dari sisi keamanan, sistem dapat ditingkatkan dengan menambahkan autentikasi tambahan seperti RFID, PIN, atau QR Code. Fitur tersebut dapat digunakan sebagai verifikasi langsung pada perangkat loker, sehingga akses tidak hanya bergantung pada web. Selain itu, keamanan komunikasi MQTT juga perlu diperhatikan dengan menggunakan koneksi TLS, username, password, dan topic yang tidak mudah ditebak."
)
para(
    "Untuk pengembangan perangkat keras, desain rangkaian dapat dibuat lebih rapi menggunakan PCB agar pemasangan komponen menjadi lebih kuat dan aman. Catu daya juga perlu diperhitungkan dengan baik agar mampu menyuplai ESP32, servo, dan LED secara stabil. Dengan pengembangan tersebut, sistem smart loker dapat bekerja lebih optimal dan siap digunakan pada lingkungan nyata."
)

doc.save(OUTPUT)
print(OUTPUT)
