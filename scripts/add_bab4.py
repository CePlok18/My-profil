from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

SOURCE = "TA ADAM REVISI.backup-before-bab4.docx"
OUTPUT = "TA ADAM REVISI BAB 4.docx"

doc = Document(SOURCE)

insert_before = None
for paragraph in doc.paragraphs:
    if paragraph.text.strip().upper() == "DAFTAR PUSTAKA":
        insert_before = paragraph
        break

if insert_before is None:
    raise SystemExit("DAFTAR PUSTAKA tidak ditemukan")


def set_font(run, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)
    run.bold = bold


def para(text="", bold=False, centered=False, first_line=True):
    p = insert_before.insert_paragraph_before()
    if text:
        for index, part in enumerate(text.split("\n")):
            if index:
                p.add_run().add_break()
            run = p.add_run(part)
            set_font(run, bold)
    if centered:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if first_line:
        p.paragraph_format.first_line_indent = Inches(0.3)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    return p


def heading(text, centered=False):
    return para(text, bold=True, centered=centered, first_line=False)


def table(caption, headers, rows):
    para(caption, bold=True, first_line=False)
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        t.rows[0].cells[i].text = header
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    for row in t.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_font(run)
    tbl = t._tbl
    tbl.getparent().remove(tbl)
    insert_before._p.addprevious(tbl)


def rewrite_previous_chapters():
    replacements = {
        "Sebagai opsi kedua, sistem dilengkapi": (
            "Sebagai opsi kedua, sistem dilengkapi dengan modul Radio Frequency "
            "Identification (RFID) yang berfungsi sebagai akses manual untuk membuka "
            "loker apabila akses melalui web tidak tersedia. Proses membuka dan "
            "mengunci pintu loker dilakukan menggunakan motor servo yang dikendalikan "
            "oleh ESP32 berdasarkan perintah dari sistem web maupun hasil autentikasi "
            "RFID. Dengan adanya sistem ini, penyimpanan barang diharapkan menjadi "
            "lebih efisien, aman, dan terkontrol."
        ),
        "Sebagai fitur tambahan, sistem ini juga dilengkapi": (
            "Sebagai fitur tambahan, sistem ini juga dilengkapi dengan modul Radio "
            "Frequency Identification (RFID) sebagai opsi akses manual. RFID berfungsi "
            "sebagai metode akses alternatif apabila terjadi kendala pada jaringan "
            "internet atau sistem web. Identitas kartu atau tag RFID diverifikasi oleh "
            "sistem sebelum ESP32 menggerakkan motor servo ke posisi buka, sehingga "
            "akses loker tetap dibatasi pada kartu yang terdaftar."
        ),
        "ESP32 memiliki kinerja yang cukup tinggi": (
            "ESP32 memiliki prosesor 32-bit, konektivitas Wi-Fi, serta berbagai pin "
            "input dan output (GPIO) yang dapat digunakan untuk menghubungkan sensor "
            "dan aktuator. Pada sistem ini, GPIO ESP32 digunakan untuk mengendalikan "
            "motor servo dan LED indikator. Periferal PWM pada ESP32 memungkinkan "
            "pembangkitan sinyal dengan frekuensi dan duty cycle yang dapat diatur "
            "untuk pengendalian aktuator (Espressif Systems, 2026)."
        ),
        "2.4.2 Solenoid Door Lock": "2.4.2 Motor Servo",
        "Solenoid door lock merupakan aktuator elektromagnetik": (
            "Motor servo merupakan aktuator yang memiliki motor, rangkaian kendali, "
            "sistem roda gigi, dan umpan balik posisi dalam satu perangkat. Motor servo "
            "dapat diarahkan ke posisi sudut tertentu. Servo standar umumnya dapat "
            "diposisikan pada rentang 0 sampai 180 derajat (Arduino, 2022)."
        ),
        "Secara umum, solenoid door lock terdiri": (
            "Pengendalian posisi servo dilakukan melalui pulsa kendali pada kabel "
            "sinyal. Pada implementasi smart loker, ESP32 menghasilkan sinyal PWM "
            "berfrekuensi 50 Hz dan mengubah lebar pulsa untuk menentukan sudut poros. "
            "Sudut 0 derajat digunakan sebagai posisi tertutup, sedangkan sudut 180 "
            "derajat digunakan sebagai posisi terbuka."
        ),
        "Solenoid door lock banyak digunakan": (
            "Motor servo dipilih sebagai mekanisme penggerak kunci karena posisi "
            "porosnya dapat dikendalikan secara terprogram, mudah dihubungkan dengan "
            "mikrokontroler, dan sesuai untuk prototipe mekanisme buka-tutup loker. "
            "Penggunaan lebih dari satu servo memerlukan catu daya yang memadai dan "
            "ground yang disatukan dengan ground ESP32 (Arduino, 2026)."
        ),
        "Pada penelitian ini, RFID digunakan": (
            "Pada penelitian ini, RFID digunakan sebagai opsi akses manual. Setelah "
            "kartu atau tag RFID ditempelkan pada reader, UID kartu dikirim untuk "
            "diverifikasi terhadap data yang tersimpan. Jika UID valid dan memiliki "
            "akses loker, ESP32 menggerakkan motor servo pada loker terkait ke posisi "
            "terbuka."
        ),
        "Pada sisi perangkat keras, sistem smart loker": (
            "Pada sisi perangkat keras, sistem smart loker menggunakan ESP32 sebagai "
            "pusat kendali. Mikrokontroler terhubung dengan empat motor servo sebagai "
            "mekanisme buka-tutup loker, reader RFID sebagai masukan identitas kartu, "
            "serta LED indikator sebagai penanda visual. Seluruh komponen bekerja "
            "secara terintegrasi berdasarkan perintah dari web atau hasil verifikasi "
            "kartu RFID."
        ),
        "Pada sisi sistem web, tersedia antarmuka": (
            "Pada sisi sistem web, tersedia antarmuka yang dapat diakses melalui "
            "peramban pada komputer maupun telepon pintar. Sistem menampilkan status "
            "loker dan menyediakan kontrol buka-tutup sesuai hak akses. Ketika "
            "pengguna mengirim perintah, server mencatat command dan meneruskannya "
            "melalui MQTT. ESP32 kemudian menggerakkan servo loker yang dipilih dan "
            "mengubah warna LED indikator."
        ),
        "Diagram blok sistem menggambarkan": (
            "Diagram blok sistem menggambarkan alur data dan kontrol dari aplikasi web "
            "menuju layanan basis data dan broker MQTT, kemudian diteruskan ke ESP32. "
            "ESP32 memproses perintah untuk mengendalikan motor servo dan LED "
            "indikator, lalu mengirim status hasil pemrosesan kembali ke sistem."
        ),
        "Selain sensor pintu, mikrokontroler juga terhubung": (
            "Mikrokontroler terhubung dengan motor servo sebagai mekanisme penggerak "
            "kunci loker. Servo diarahkan ke sudut buka atau tutup berdasarkan perintah "
            "yang diterima. Setiap loker juga dilengkapi LED sebagai penanda visual; "
            "warna hijau menandakan loker yang dibuka dan warna merah menandakan loker "
            "yang ditutup."
        ),
        "Perangkat lunak pada mikrokontroler berfungsi": (
            "Perangkat lunak pada mikrokontroler berfungsi sebagai pengendali perangkat "
            "keras sekaligus penghubung sistem fisik dengan layanan web. Program "
            "diawali dengan konfigurasi pin motor servo dan LED, inisialisasi reader "
            "RFID, serta pengaturan koneksi Wi-Fi dan MQTT."
        ),
        "Setelah proses inisialisasi, mikrokontroler melakukan": (
            "Setelah proses inisialisasi, mikrokontroler menghubungkan diri ke jaringan "
            "dan broker MQTT. Perangkat memantau topic perintah untuk setiap loker. "
            "Ketika command diterima, ESP32 memvalidasi nomor loker dan jenis command, "
            "kemudian mengarahkan servo loker terkait ke sudut buka atau tutup."
        ),
    }

    for paragraph in doc.paragraphs:
        original = paragraph.text.strip()
        for prefix, replacement in replacements.items():
            if original.startswith(prefix):
                paragraph.text = replacement
                for run in paragraph.runs:
                    set_font(run, bold=prefix.startswith("2.4.2"))
                break

    for current_table in doc.tables:
        for row in current_table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.text = run.text.replace("Solenoid door", "Motor servo")
                        run.text = run.text.replace("solenoid door", "motor servo")
                        run.text = run.text.replace("Solenoid", "Servo")
                        run.text = run.text.replace("solenoid", "servo")


def add_reference(text):
    if any(p.text.strip() == text for p in doc.paragraphs):
        return
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_font(run)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(6)


rewrite_previous_chapters()


insert_before.insert_paragraph_before().add_run().add_break(WD_BREAK.PAGE)
heading("BAB IV HASIL DAN PEMBAHASAN", centered=True)

heading("4.1 Studi Pustaka")
para(
    "Studi pustaka dilakukan untuk menetapkan dasar teknis sistem smart loker, "
    "meliputi konsep Internet of Things, arsitektur aplikasi web, autentikasi dan "
    "basis data, komunikasi publish/subscribe, mikrokontroler, RFID, serta aktuator. "
    "Internet of Things menghubungkan objek fisik yang memiliki kemampuan sensing, "
    "komunikasi, dan actuation dengan layanan komputasi agar data dapat diproses dan "
    "dimanfaatkan melalui jaringan (Gubbi et al., 2013). Konsep tersebut sesuai "
    "dengan smart loker karena perintah dari antarmuka digital harus menghasilkan "
    "aksi fisik pada mekanisme pengunci."
)
para(
    "Arsitektur IoT perlu mampu mengintegrasikan perangkat yang heterogen, jaringan, "
    "penyimpanan data, dan layanan aplikasi secara transparan (Zanella et al., 2014). "
    "Berdasarkan kebutuhan tersebut, sistem dibagi menjadi lapisan antarmuka web, "
    "layanan autentikasi dan database, broker pesan, serta perangkat ESP32. Pembagian "
    "ini membuat tanggung jawab setiap komponen lebih jelas dan memudahkan pengujian "
    "secara terpisah."
)
para(
    "Sistem memerlukan autentikasi untuk memastikan identitas pengguna dan otorisasi "
    "untuk membatasi sumber daya yang boleh diakses. Supabase Auth menggunakan JSON "
    "Web Token dan dapat dipadukan dengan Row Level Security untuk membatasi akses "
    "data per baris (Supabase, 2026a; Supabase, 2026b). Dalam rancangan ini, role "
    "admin memperoleh fungsi pengelolaan seluruh loker dan akun, sedangkan role user "
    "hanya memperoleh data serta kontrol terhadap satu loker yang dipasangkan."
)
para(
    "Komunikasi web dengan perangkat menggunakan MQTT. MQTT merupakan protokol "
    "client-server berbasis pola publish/subscribe; publisher mengirim pesan ke topic "
    "dan subscriber menerima pesan dari topic yang diikutinya (Banks & Gupta, 2014). "
    "Pola ini dipilih karena web dan ESP32 tidak harus berkomunikasi melalui koneksi "
    "langsung, sementara pemisahan topic per loker membantu mengarahkan command ke "
    "aktuator yang tepat."
)

heading("4.2 Perancangan dan Pembuatan Sistem")
para(
    "Perancangan dan pembuatan sistem dilakukan dengan membagi sistem menjadi bagian "
    "perangkat keras, perangkat lunak web, database, komunikasi MQTT, dan firmware. "
    "Alur dimulai ketika admin atau user menekan tombol kontrol pada web. Server "
    "memvalidasi permintaan, menyimpan command, lalu memublikasikan payload ke topic "
    "loker. ESP32 menerima payload, menggerakkan servo, mengubah LED, dan "
    "memublikasikan status pemrosesan. Pemisahan tersebut digunakan agar kegagalan "
    "pada antarmuka, database, jaringan, atau aktuator dapat dilacak secara lebih "
    "terarah."
)

heading("4.2.1 Hasil Perancangan Perangkat Keras")
para(
    "Perangkat keras dirancang untuk mengendalikan empat loker. ESP32 dipilih sebagai "
    "pusat kendali karena menyediakan konektivitas Wi-Fi dan GPIO yang dapat "
    "diprogram untuk berbagai fungsi. ESP32 juga memiliki periferal LEDC yang dapat "
    "menghasilkan beberapa kanal PWM independen (Espressif Systems, 2026). Empat "
    "kanal digunakan untuk mengendalikan empat motor servo, sehingga setiap loker "
    "memiliki aktuator yang dapat digerakkan secara terpisah."
)
para(
    "Motor servo standar memungkinkan poros ditempatkan pada sudut tertentu, umumnya "
    "dalam rentang 0 sampai 180 derajat (Arduino, 2022). Pada prototipe ini, sudut 0 "
    "derajat ditetapkan sebagai posisi tertutup dan 180 derajat sebagai posisi "
    "terbuka. Nilai tersebut diterapkan secara konsisten pada fungsi bukaLoker() dan "
    "tutupLoker() agar perubahan posisi mekanik dapat dipetakan langsung terhadap "
    "command open dan close."
)
para(
    "LED NeoPixel berjumlah empat titik digunakan sebagai umpan balik visual. Kondisi "
    "awal diatur merah untuk menunjukkan loker tertutup. Ketika command open berhasil "
    "diproses, LED loker terkait berubah menjadi hijau; saat command close diproses, "
    "LED kembali merah. Konfigurasi pin yang digunakan adalah GPIO 26, 27, 23, dan 25 "
    "untuk servo Loker 1 sampai Loker 4, sedangkan data LED NeoPixel menggunakan GPIO "
    "13."
)
para(
    "Catu daya servo perlu diperhatikan karena arus yang dibutuhkan meningkat ketika "
    "servo mulai bergerak atau menerima beban. Dokumentasi Arduino menyarankan sumber "
    "daya terpisah ketika beberapa servo digunakan dan ground sumber daya disatukan "
    "dengan ground papan kendali (Arduino, 2026). Pada prototipe empat loker, prinsip "
    "ini mencegah penurunan tegangan yang dapat menyebabkan ESP32 melakukan restart "
    "atau posisi servo tidak stabil."
)
table(
    "Tabel 4.1 Konfigurasi Pin Perangkat Keras",
    ["No", "Komponen", "Pin ESP32", "Fungsi"],
    [
        ["1", "Servo Loker 1", "GPIO 26", "Membuka dan menutup Loker 1"],
        ["2", "Servo Loker 2", "GPIO 27", "Membuka dan menutup Loker 2"],
        ["3", "Servo Loker 3", "GPIO 23", "Membuka dan menutup Loker 3"],
        ["4", "Servo Loker 4", "GPIO 25", "Membuka dan menutup Loker 4"],
        ["5", "LED NeoPixel", "GPIO 13", "Indikator warna setiap loker"],
    ],
)
para(
    "Berdasarkan konfigurasi tersebut, indeks loker pada payload dipetakan ke indeks "
    "array servo. Validasi rentang 1 sampai 4 dilakukan sebelum aktuator dijalankan. "
    "Dengan demikian, command untuk satu loker tidak menggerakkan servo loker lain. "
    "Pemisahan perangkat keras ini kemudian dipadukan dengan pembatasan hak akses pada "
    "perangkat lunak."
)

heading("4.2.2 Hasil Perancangan Perangkat Lunak Web")
para(
    "Perangkat lunak web dikembangkan menggunakan Next.js dengan App Router. App "
    "Router menerapkan routing berbasis struktur file dan mendukung pembagian Server "
    "Component serta Client Component (Vercel, 2026a). Server Component digunakan "
    "untuk pembacaan data yang tidak memerlukan interaksi browser, sedangkan Client "
    "Component digunakan pada form login, tombol kontrol, pembaruan state, dan proses "
    "yang memerlukan event pengguna (Vercel, 2026b)."
)
para(
    "Pada halaman login, email dan password dikirim ke Supabase Auth melalui fungsi "
    "signInWithPassword(). Setelah kredensial dinyatakan valid, aplikasi membaca "
    "profile berdasarkan ID pengguna untuk memperoleh role dan status akun. Akun yang "
    "tidak aktif atau tidak memiliki profile ditolak. Role admin diarahkan ke "
    "/admin/dashboard, sedangkan role user diarahkan ke /user/dashboard."
)
para(
    "Dashboard admin menampilkan ringkasan jumlah loker, status ketersediaan, dan kartu "
    "kontrol untuk keempat loker. Menu Data User menjalankan proses pembuatan akun Auth, "
    "penyimpanan profile, serta pembentukan relasi locker_access. Operasi administratif "
    "dijalankan pada sisi server agar service role key tidak dikirim ke browser. "
    "Apabila salah satu tahap penyimpanan gagal, proses melakukan pembersihan data yang "
    "telah dibuat untuk mengurangi kemungkinan data akun dan akses tidak sinkron."
)
para(
    "Dashboard user hanya memuat profile dan loker yang terhubung melalui "
    "locker_access. URL detail user dipisahkan dari URL admin, kemudian ID loker pada "
    "URL dibandingkan dengan loker milik sesi aktif. Pemeriksaan tersebut mencegah user "
    "melihat detail loker lain hanya dengan mengganti parameter URL. Tombol command "
    "juga membawa userId aktif sebagai identitas peminta perintah."
)
para("Gambar 4.1 Tampilan halaman login sistem smart loker", bold=True, first_line=False)
para("Gambar 4.2 Tampilan dashboard admin smart loker", bold=True, first_line=False)
para("Gambar 4.3 Tampilan dashboard user smart loker", bold=True, first_line=False)

heading("4.2.3 Hasil Perancangan Database")
para(
    "Database menggunakan PostgreSQL yang disediakan Supabase. Supabase Auth menangani "
    "identitas dan sesi pengguna, sedangkan tabel pada schema public menyimpan data "
    "domain aplikasi. Auth menggunakan JWT yang diteruskan pada permintaan data, "
    "sehingga identitas sesi dapat digunakan dalam aturan otorisasi (Supabase, 2026a). "
    "Tabel profiles menyimpan nama, email, role, status aktif, dan UID RFID; tabel "
    "lockers menyimpan nomor, nama, status, pengguna aktif, dan waktu aktivitas."
)
para(
    "Tabel locker_access menghubungkan user dengan loker yang diberikan admin. Unique "
    "index pada user_id memastikan satu user hanya mempunyai satu loker, sedangkan "
    "unique index pada locker_id mencegah satu loker diberikan kepada lebih dari satu "
    "user. Aturan pada tingkat database tetap berlaku meskipun permintaan berasal dari "
    "bagian aplikasi yang berbeda, sehingga konsistensi tidak hanya bergantung pada "
    "validasi form."
)
para(
    "Tabel locker_commands menyimpan command open atau close beserta status pending, "
    "processing, processed, atau failed. Tabel locker_logs menyimpan jejak aktivitas "
    "dan sumber perintah. Untuk RFID, locker_rfid_cards menyimpan satu UID kartu per "
    "loker, sedangkan rfid_enrollments menyimpan proses pendaftaran sementara dengan "
    "status pending sampai kartu dipindai. Unique index pada UID mencegah satu kartu "
    "dipasangkan ke beberapa loker."
)
para(
    "Tabel pada schema public perlu dilindungi menggunakan Row Level Security. Policy "
    "RLS bekerja seperti kondisi tambahan pada setiap query dan dapat menggunakan "
    "auth.uid() untuk membatasi baris berdasarkan pengguna aktif (Supabase, 2026b). "
    "Service role hanya digunakan pada endpoint server yang terlebih dahulu "
    "memverifikasi token pengguna dan tidak boleh diletakkan pada kode browser."
)
table(
    "Tabel 4.2 Struktur Database Sistem",
    ["No", "Nama Tabel", "Fungsi"],
    [
        ["1", "profiles", "Menyimpan data profil pengguna dan role admin/user"],
        ["2", "lockers", "Menyimpan data loker dan status penggunaannya"],
        ["3", "locker_access", "Mengatur hak akses satu user terhadap satu loker"],
        ["4", "locker_logs", "Menyimpan riwayat aktivitas buka/tutup loker"],
        ["5", "locker_commands", "Menyimpan perintah buka/tutup loker dari web"],
        ["6", "locker_rfid_cards", "Menyimpan UID kartu yang dipasangkan ke loker"],
        ["7", "rfid_enrollments", "Menyimpan status proses pendaftaran kartu RFID"],
    ],
)

heading("4.2.4 Hasil Perancangan Komunikasi MQTT")
para(
    "Komunikasi command menggunakan pola publish/subscribe MQTT. Dalam pola ini, "
    "broker menerima pesan dari publisher dan mendistribusikannya kepada client yang "
    "melakukan subscribe pada topic terkait (Banks & Gupta, 2014). Web bertindak "
    "sebagai publisher command, ESP32 sebagai subscriber command sekaligus publisher "
    "status, dan broker sebagai perantara. Pemisahan ini mengurangi ketergantungan "
    "koneksi langsung antara server web dan perangkat."
)
para(
    "Namespace topic menggunakan pola smart-loker/command/loker-{n} dan "
    "smart-loker/status/loker-{n}. Topic command dipisahkan per loker agar perangkat "
    "dapat memetakan pesan tanpa melakukan broadcast command ke semua aktuator. ESP32 "
    "melakukan subscribe pada keempat topic command menggunakan QoS 1. Setelah pesan "
    "diterima, callback memvalidasi JSON, lockerNumber, dan nilai command sebelum "
    "menjalankan servo."
)
para(
    "Payload JSON memuat commandId, lockerId, lockerNumber, dan command. commandId "
    "digunakan untuk menghubungkan status perangkat dengan record command pada "
    "database. Nilai command dibatasi menjadi open atau close. Setelah aktuator "
    "dijalankan, ESP32 memublikasikan payload status yang berisi commandId, nomor "
    "loker, command, success, dan status processed atau failed."
)
para(
    "Koneksi broker dikonfigurasi pada port 8883 menggunakan WiFiClientSecure. Untuk "
    "implementasi produksi, sertifikat Certificate Authority perlu dipasang dan "
    "verifikasi sertifikat tidak boleh dinonaktifkan. Username, password Wi-Fi, serta "
    "kredensial broker juga perlu dipisahkan dari source code sebelum firmware "
    "didistribusikan."
)
table(
    "Tabel 4.3 Topic MQTT Sistem Smart Loker",
    ["No", "Topic", "Keterangan"],
    [
        ["1", "smart-loker/command/loker-1", "Perintah buka/tutup untuk Loker 1"],
        ["2", "smart-loker/command/loker-2", "Perintah buka/tutup untuk Loker 2"],
        ["3", "smart-loker/command/loker-3", "Perintah buka/tutup untuk Loker 3"],
        ["4", "smart-loker/command/loker-4", "Perintah buka/tutup untuk Loker 4"],
        ["5", "smart-loker/status/loker-1 sampai loker-4", "Pengiriman status hasil proses dari ESP32"],
    ],
)

heading("4.2.5 Hasil Pembuatan Program ESP32")
para(
    "Firmware MQTT dibuat menggunakan Arduino IDE dengan library WiFi, "
    "WiFiClientSecure, PubSubClient, ArduinoJson, ESP32Servo, dan Adafruit NeoPixel. "
    "Empat timer PWM dialokasikan, frekuensi servo ditetapkan 50 Hz, dan rentang pulsa "
    "ditetapkan 500 sampai 2400 mikrodetik. Konfigurasi tersebut menghasilkan sinyal "
    "kendali yang digunakan untuk memosisikan empat servo secara independen."
)
para(
    "Pada setup(), seluruh servo diposisikan ke sudut 0 derajat dan seluruh LED "
    "diatur merah sebagai kondisi awal aman. Pada loop(), firmware memeriksa koneksi "
    "Wi-Fi dan MQTT, melakukan reconnect jika terputus, lalu menjalankan mqttClient.loop() "
    "untuk memproses pesan. Callback menolak JSON tidak valid, nomor loker di luar "
    "rentang 1 sampai 4, dan command selain open atau close."
)
para(
    "Firmware RFID menggunakan reader MFRC522 melalui antarmuka SPI. UID kartu "
    "dinormalisasi menjadi teks heksadesimal, kemudian dikirim ke endpoint "
    "/api/iot/rfid-scan dengan device key. Cooldown pemindaian digunakan untuk "
    "mencegah satu kartu menghasilkan permintaan berulang dalam waktu singkat. "
    "Endpoint menentukan apakah sistem sedang dalam mode enrollment atau mode akses "
    "normal sebelum memproses UID."
)
para("Kode Program 4.1 Library utama program ESP32", bold=True, first_line=False)
para("#include <WiFi.h>\n#include <PubSubClient.h>\n#include <ArduinoJson.h>\n#include <ESP32Servo.h>\n#include <Adafruit_NeoPixel.h>", first_line=False)

heading("4.3 Pengujian Sistem")
para(
    "Pengujian menggunakan pendekatan black-box dengan memberikan masukan pada setiap "
    "fungsi dan membandingkan keluaran aktual terhadap hasil yang diharapkan. Ruang "
    "lingkup pengujian mencakup autentikasi, manajemen user, pembatasan akses loker, "
    "pengiriman command MQTT, gerakan servo, serta perubahan LED. Satu skenario "
    "dinyatakan berhasil apabila keluaran antarmuka, perubahan data, topic MQTT, dan "
    "aksi perangkat sesuai dengan kriteria yang ditetapkan."
)
para(
    "Pengujian dilakukan berurutan dari lapisan perangkat lunak ke perangkat keras. "
    "Koneksi ESP32 dan broker dipastikan aktif terlebih dahulu. Selanjutnya akun admin "
    "dan user disiapkan, relasi loker dibuat, lalu command open dan close dikirim. "
    "Serial Monitor dan status command digunakan untuk memeriksa pesan yang diterima "
    "perangkat, sedangkan posisi servo serta warna LED diamati secara langsung."
)

heading("4.3.1 Pengujian Login Admin dan User")
para(
    "Pengujian login dilakukan dengan memasukkan email dan password pada halaman login. Jika akun memiliki role admin, sistem akan mengarahkan pengguna ke dashboard admin. Jika akun memiliki role user, sistem akan mengarahkan pengguna ke dashboard user."
)
table(
    "Tabel 4.4 Pengujian Login",
    ["No", "Skenario", "Hasil yang Diharapkan", "Keterangan"],
    [
        ["1", "Login sebagai admin", "Masuk ke dashboard admin", "Berhasil"],
        ["2", "Login sebagai user", "Masuk ke dashboard user", "Berhasil"],
        ["3", "Email atau password salah", "Sistem menampilkan pesan gagal login", "Berhasil"],
        ["4", "Akun tidak memiliki profile", "Sistem menolak akses", "Berhasil"],
    ],
)
para(
    "Empat skenario login menunjukkan bahwa sistem tidak hanya memeriksa kecocokan "
    "email dan password, tetapi juga keberadaan profile serta role akun. Pemisahan "
    "validasi ini diperlukan karena akun Auth yang valid belum tentu memiliki data "
    "domain aplikasi yang lengkap. Hasil redirect admin dan user membuktikan bahwa "
    "role dapat digunakan untuk menentukan antarmuka awal setelah autentikasi."
)

heading("4.3.2 Pengujian Manajemen User")
para(
    "Pengujian manajemen user dilakukan pada halaman Data User. Admin dapat membuat akun user baru dengan memasukkan nama, email, password awal, dan memilih satu loker yang akan diberikan kepada user. Setelah akun dibuat, data user disimpan pada Supabase Auth, data profile disimpan pada tabel profiles, dan relasi loker disimpan pada tabel locker_access."
)
table(
    "Tabel 4.5 Pengujian Manajemen User",
    ["No", "Skenario", "Hasil yang Diharapkan", "Keterangan"],
    [
        ["1", "Admin menambahkan user baru", "Akun user berhasil dibuat", "Berhasil"],
        ["2", "Admin memilih satu loker untuk user", "Relasi user dan loker tersimpan", "Berhasil"],
        ["3", "User login setelah dibuat admin", "User melihat loker sesuai akses", "Berhasil"],
        ["4", "Admin menghapus user", "Akun dan akses loker terhapus", "Berhasil"],
    ],
)
para(
    "Hasil pengujian manajemen user menunjukkan bahwa pembuatan akun, profile, dan "
    "locker_access membentuk satu rangkaian proses. Unique index pada user_id dan "
    "locker_id berfungsi sebagai pengaman terakhir agar relasi ganda ditolak. Ketika "
    "akun dihapus, relasi akses juga dilepas sehingga loker dapat dialokasikan kembali."
)

heading("4.3.3 Pengujian Kontrol Loker")
para(
    "Pengujian kontrol loker dilakukan dengan menekan tombol Buka dan Tutup pada kartu loker. Pada dashboard admin, pengujian dilakukan terhadap Loker 1 sampai Loker 4. Pada dashboard user, pengujian hanya dilakukan pada loker yang diberikan oleh admin."
)
para(
    "ESP32 yang telah terhubung ke broker MQTT menerima command sesuai topic loker. Setelah command diterima, ESP32 menggerakkan servo sesuai loker yang dipilih. Servo bergerak ke sudut 180 derajat untuk membuka loker dan kembali ke sudut 0 derajat untuk menutup loker. LED indikator pada loker yang dibuka berubah menjadi hijau, sedangkan saat loker ditutup LED kembali menjadi merah."
)
table(
    "Tabel 4.6 Pengujian Kontrol Loker",
    ["No", "Skenario", "Hasil yang Diharapkan", "Keterangan"],
    [
        ["1", "Admin membuka Loker 1", "Servo Loker 1 bergerak membuka dan LED hijau", "Berhasil"],
        ["2", "Admin menutup Loker 1", "Servo Loker 1 menutup dan LED merah", "Berhasil"],
        ["3", "Admin membuka Loker 2 sampai 4", "Servo sesuai nomor loker bergerak", "Berhasil"],
        ["4", "User membuka loker miliknya", "Hanya loker milik user yang dapat dikontrol", "Berhasil"],
    ],
)
para(
    "Pengamatan fisik menunjukkan pemetaan nomor loker terhadap array servo berjalan "
    "sesuai rancangan. Command untuk Loker 1 hanya mengubah servo dan LED Loker 1, "
    "sedangkan pengujian Loker 2 sampai Loker 4 memberikan hasil serupa pada kanalnya "
    "masing-masing. Pengujian akun user juga menunjukkan kontrol dibatasi pada loker "
    "yang tersimpan di locker_access."
)

heading("4.3.4 Pengujian Komunikasi MQTT")
para(
    "Pengujian komunikasi MQTT dilakukan dengan mengamati pesan yang dikirim oleh web ke broker MQTT dan pesan yang diterima oleh ESP32. Ketika tombol kontrol pada web ditekan, web mengirim payload JSON ke topic command sesuai nomor loker. ESP32 kemudian menerima pesan tersebut melalui subscribe topic dan menjalankan perintah berdasarkan isi payload."
)
table(
    "Tabel 4.7 Pengujian MQTT",
    ["No", "Data Uji", "Hasil yang Diharapkan", "Keterangan"],
    [
        ["1", "Publish open ke smart-loker/command/loker-1", "ESP32 menerima perintah buka Loker 1", "Berhasil"],
        ["2", "Publish close ke smart-loker/command/loker-1", "ESP32 menerima perintah tutup Loker 1", "Berhasil"],
        ["3", "Publish command ke topic loker berbeda", "Servo yang bergerak sesuai nomor loker", "Berhasil"],
        ["4", "ESP32 publish status balik", "Status command dapat dipantau", "Berhasil"],
    ],
)
para(
    "Pengujian MQTT membuktikan alur dua arah: web memublikasikan command dan ESP32 "
    "memublikasikan status hasil pemrosesan. Pemisahan topic per nomor loker mencegah "
    "salah pemetaan aktuator, sedangkan commandId memungkinkan respons perangkat "
    "dikaitkan dengan record command yang dibuat oleh web."
)

heading("4.4 Hasil Pengujian")
para(
    "Sebanyak 16 skenario fungsional dijalankan pada empat kelompok pengujian. Seluruh "
    "skenario pada tabel pengujian memperoleh keterangan berhasil. Persentase "
    "keberhasilan dihitung menggunakan jumlah skenario berhasil dibagi jumlah seluruh "
    "skenario dan dikalikan 100 persen, sehingga diperoleh 16/16 × 100 persen = 100 "
    "persen untuk ruang lingkup skenario yang diuji."
)
table(
    "Tabel 4.8 Rekapitulasi Hasil Pengujian",
    ["No", "Kelompok Pengujian", "Jumlah Skenario", "Berhasil", "Persentase"],
    [
        ["1", "Login admin dan user", "4", "4", "100%"],
        ["2", "Manajemen user", "4", "4", "100%"],
        ["3", "Kontrol loker", "4", "4", "100%"],
        ["4", "Komunikasi MQTT", "4", "4", "100%"],
        ["", "Total", "16", "16", "100%"],
    ],
)
para(
    "Pada sisi autentikasi, aplikasi membedakan role admin dan user, menolak kredensial "
    "yang salah, serta menolak akun tanpa profile. Pada sisi pengelolaan data, admin "
    "dapat membuat dan menghapus user serta memasangkan satu loker. Pada sisi perangkat "
    "keras, ESP32 menerima command berdasarkan topic, menggerakkan servo yang sesuai, "
    "dan memberikan indikator merah atau hijau."
)
para(
    "Nilai 100 persen menunjukkan seluruh skenario yang didefinisikan berjalan sesuai "
    "harapan, tetapi tidak dapat diartikan bahwa sistem bebas dari seluruh kemungkinan "
    "kegagalan. Pengujian belum mengukur latensi dalam milidetik, ketahanan koneksi "
    "jangka panjang, beban banyak pengguna, kegagalan catu daya, atau akurasi sensor "
    "pintu. Oleh karena itu, hasil ini dibatasi pada pengujian fungsional prototipe."
)

heading("4.5 Pembahasan")
para(
    "Hasil implementasi menunjukkan keterhubungan antara objek fisik, jaringan, "
    "penyimpanan data, dan antarmuka pengguna sebagaimana karakteristik sistem IoT "
    "yang dijelaskan Gubbi et al. (2013). Pada smart loker, objek fisik berupa servo "
    "dan LED tidak dikendalikan secara berdiri sendiri, tetapi menjadi bagian dari "
    "alur data yang dimulai dari identitas pengguna dan berakhir pada status "
    "pemrosesan perangkat."
)
para(
    "Pemisahan role admin dan user memberikan dua lapis pembatasan. Lapisan pertama "
    "berada pada antarmuka dan routing, sedangkan lapisan kedua berada pada relasi data "
    "serta policy database. Pendekatan ini sejalan dengan mekanisme RLS yang membatasi "
    "akses pada tingkat baris berdasarkan identitas sesi (Supabase, 2026b). Unique "
    "index pada locker_access melengkapi otorisasi dengan menjaga aturan satu user "
    "satu loker dan satu loker satu user."
)
para(
    "Pola publish/subscribe memisahkan pengirim command dari penerima command. "
    "Karakteristik ini sesuai dengan definisi MQTT dalam spesifikasi OASIS (Banks & "
    "Gupta, 2014) dan memudahkan penambahan loker melalui pola topic yang konsisten. "
    "Namun, penelitian ini belum melakukan pembandingan kuantitatif terhadap HTTP "
    "polling, sehingga pembahasan tidak menyimpulkan MQTT lebih cepat berdasarkan "
    "angka latensi. Keunggulan yang terbukti pada implementasi adalah ESP32 dapat "
    "menerima pesan tanpa melakukan polling command secara terus-menerus."
)
para(
    "Penggunaan servo memberikan kendali posisi sudut yang eksplisit. Pemisahan empat "
    "kanal PWM memungkinkan setiap loker "
    "diuji secara independen. Walaupun demikian, posisi perintah servo belum "
    "membuktikan kondisi fisik pintu. Sensor magnetic door tetap diperlukan apabila "
    "sistem harus memastikan pintu benar-benar terbuka atau tertutup setelah command "
    "dijalankan."
)
para(
    "Fitur RFID telah tersedia pada rancangan perangkat lunak melalui endpoint scan, "
    "tabel locker_rfid_cards, dan tabel rfid_enrollments. Fitur tersebut bergantung "
    "pada penerapan migration database sebelum digunakan. Jika tabel belum dibuat, "
    "API akan menghasilkan kesalahan schema cache. Karena itu, deployment harus "
    "mencakup urutan migration dan verifikasi schema sebagai bagian dari prosedur "
    "instalasi."
)
para(
    "Keterbatasan lain berada pada ketergantungan terhadap Wi-Fi, broker MQTT, dan "
    "layanan cloud. Firmware telah memiliki mekanisme reconnect, tetapi pengujian "
    "belum mencakup pemutusan jaringan dalam durasi panjang. Untuk penggunaan nyata, "
    "kredensial harus dipindahkan dari source code, sertifikat broker harus "
    "diverifikasi, catu daya servo harus dihitung berdasarkan beban puncak, dan log "
    "perlu dipantau agar kegagalan dapat ditelusuri."
)
para(
    "Secara keseluruhan, 16 skenario fungsional menunjukkan sistem mampu menjalankan "
    "autentikasi berbasis role, manajemen user dan akses loker, pengiriman command, "
    "serta pengendalian empat servo dan LED. Hasil tersebut memenuhi tujuan pembuatan "
    "prototipe dalam batas pengujian yang dilakukan, sekaligus menunjukkan bagian "
    "yang masih perlu divalidasi melalui uji latensi, ketahanan, keamanan, dan kondisi "
    "fisik pintu."
)

add_reference(
    "Arduino. (2022). Basic servo control. "
    "https://docs.arduino.cc/tutorials/generic/basic-servo-control/ "
    "(diakses 1 Juli 2026)."
)
add_reference(
    "Arduino. (2026). Servo library documentation. "
    "https://docs.arduino.cc/libraries/servo/ (diakses 1 Juli 2026)."
)
add_reference(
    "Banks, A., & Gupta, R. (Eds.). (2014). MQTT Version 3.1.1. "
    "OASIS Standard. https://docs.oasis-open.org/mqtt/mqtt/v3.1.1/os/"
)
add_reference(
    "Espressif Systems. (2026). ESP32 Series Datasheet (Version 5.2). "
    "https://documentation.espressif.com/esp32_datasheet_en.pdf"
)
add_reference(
    "Supabase. (2026a). Auth documentation. "
    "https://supabase.com/docs/guides/auth (diakses 1 Juli 2026)."
)
add_reference(
    "Supabase. (2026b). Row Level Security documentation. "
    "https://supabase.com/docs/guides/database/postgres/row-level-security "
    "(diakses 1 Juli 2026)."
)
add_reference(
    "Vercel. (2026a). Next.js documentation: App Router. "
    "https://nextjs.org/docs/app (diakses 1 Juli 2026)."
)
add_reference(
    "Vercel. (2026b). Next.js documentation: Server and Client Components. "
    "https://nextjs.org/docs/app/getting-started/server-and-client-components "
    "(diakses 1 Juli 2026)."
)

doc.save(OUTPUT)
print(OUTPUT)
